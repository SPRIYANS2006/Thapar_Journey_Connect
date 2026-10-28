import asyncio
import os
from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

SCREENSHOT_DIR = r"C:\Users\arora\.gemini\antigravity\scratch\thapar-journey-connect\screenshots"
os.makedirs(SCREENSHOT_DIR, exist_ok=True)

async def capture_screenshots():
    print("Launching headless browser for screenshots...")
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        context = await browser.new_context(viewport={'width': 1440, 'height': 900})
        page = await context.new_page()

        # 1. Auth Page
        await page.goto("http://localhost:5173/")
        await page.wait_for_timeout(1000)
        path_auth = os.path.join(SCREENSHOT_DIR, "1_auth_page.png")
        await page.screenshot(path=path_auth)
        print("Captured 1_auth_page.png")

        # 2. OTP Screen
        await page.fill('input[type="email"]', "student@thapar.edu")
        await page.fill('input[type="password"]', "password123")
        await page.click('button[type="submit"]')
        await page.wait_for_timeout(1000)
        path_otp = os.path.join(SCREENSHOT_DIR, "2_otp_screen.png")
        await page.screenshot(path=path_otp)
        print("Captured 2_otp_screen.png")

        # 3. Dashboard
        await page.click('button:has-text("Verify OTP & Access Dashboard")')
        await page.wait_for_timeout(1000)
        path_dash = os.path.join(SCREENSHOT_DIR, "3_dashboard.png")
        await page.screenshot(path=path_dash)
        print("Captured 3_dashboard.png")

        # 4. Journey Form Modal
        buttons = page.locator('button:has-text("Find Travel Partners")')
        await buttons.nth(0).click()
        await page.wait_for_timeout(1000)
        path_form = os.path.join(SCREENSHOT_DIR, "4_journey_form.png")
        await page.screenshot(path=path_form)
        print("Captured 4_journey_form.png")

        # 5. Matching Results
        modal_submit = page.locator('form button:has-text("Find Travel Partners")')
        await modal_submit.click()
        await page.wait_for_timeout(1000)
        path_match = os.path.join(SCREENSHOT_DIR, "5_matching_results.png")
        await page.screenshot(path=path_match)
        print("Captured 5_matching_results.png")

        # 6. Group Details (3/4 members)
        grp_btn = page.locator('button:has-text("View My Group & Chat"), button:has-text("Group")').first
        await grp_btn.click()
        await page.wait_for_timeout(1000)
        path_grp3 = os.path.join(SCREENSHOT_DIR, "6_group_details_3.png")
        await page.screenshot(path=path_grp3)
        print("Captured 6_group_details_3.png")

        # 7. Dynamic Join Simulator (4/4 members)
        sim_btn = page.locator('button:has-text("Simulate Student Joining")')
        await sim_btn.click()
        await page.wait_for_timeout(1000)
        path_grp4 = os.path.join(SCREENSHOT_DIR, "7_group_details_4.png")
        await page.screenshot(path=path_grp4)
        print("Captured 7_group_details_4.png")

        # 8. Group Chat Modal
        chat_btn = page.locator('button:has-text("Open Group Chat")')
        await chat_btn.click()
        await page.wait_for_timeout(1000)
        path_chat = os.path.join(SCREENSHOT_DIR, "8_group_chat.png")
        await page.screenshot(path=path_chat)
        print("Captured 8_group_chat.png")

        # Close chat modal
        close_btn = page.locator('div.fixed button').first
        await close_btn.click()
        await page.wait_for_timeout(1000)

        # 9. My Journeys Page
        await page.click('button:has-text("My Journeys")')
        await page.wait_for_timeout(1000)
        path_journeys = os.path.join(SCREENSHOT_DIR, "9_my_journeys.png")
        await page.screenshot(path=path_journeys)
        print("Captured 9_my_journeys.png")

        # 10. Profile Page
        await page.click('button:has-text("Profile")')
        await page.wait_for_timeout(1000)
        path_profile = os.path.join(SCREENSHOT_DIR, "10_profile.png")
        await page.screenshot(path=path_profile)
        print("Captured 10_profile.png")

        await browser.close()

def build_powerpoint():
    print("Building 10-slide PowerPoint presentation with enlarged text...")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    NAVY = RGBColor(15, 23, 42)
    BLUE = RGBColor(37, 99, 235)
    TEXT_DARK = RGBColor(30, 41, 59)
    ACCENT_BG = RGBColor(241, 245, 249)

    slides_data = [
        {
            "slide_num": "01",
            "tag": "EXECUTIVE OVERVIEW",
            "title": "Thapar Journey Connect",
            "subtitle": "Student Journey-Matching & Group Coordination Platform for TIET",
            "bullets": [
                "Exclusive verified platform for Thapar Institute (TIET) students.",
                "Enables shared transit to airports, railway stations, & bus stands.",
                "Reduces travel costs, improves campus safety, & cuts carbon footprint.",
                "Platform Disclaimer: Student coordination only (NOT a taxi booking app)."
            ],
            "image": os.path.join(SCREENSHOT_DIR, "1_auth_page.png")
        },
        {
            "slide_num": "02",
            "tag": "VERIFIED AUTHENTICATION",
            "title": "TIET Student Access & OTP Security",
            "subtitle": "Strict @thapar.edu domain enforcement for campus safety",
            "bullets": [
                "Email domain restriction explicitly enforces @thapar.edu addresses.",
                "Real-time validation blocks external/unauthorized email attempts.",
                "Simulated 4-Digit OTP authentication workflow for student verification.",
                "Maintains verified student profiles with hostel and contact details."
            ],
            "image": os.path.join(SCREENSHOT_DIR, "2_otp_screen.png")
        },
        {
            "slide_num": "03",
            "tag": "CENTRAL DASHBOARD",
            "title": "Transit Hubs & Core Dashboard",
            "subtitle": "Intuitive hub selection for Patiala & nearby transit points",
            "bullets": [
                "Patiala Bus Stand 🚌 - Intercity buses & local transit.",
                "Patiala Railway Station 🚆 - High-frequency train departures.",
                "Chandigarh Airport ✈️ - Shared flight cabs (Shaheed Bhagat Singh).",
                "Delhi Airport ✈️ - Long-distance cabs & train connections to IGI."
            ],
            "image": os.path.join(SCREENSHOT_DIR, "3_dashboard.png")
        },
        {
            "slide_num": "04",
            "tag": "JOURNEY CREATION",
            "title": "Intelligent Journey Window Form",
            "subtitle": "Flexible departure windows & TIET hostel block integration",
            "bullets": [
                "Departure Time Selection with ±10, ±15, ±20, ±30 minute flexibility.",
                "Supports all TIET hostels (H1-H15, Hostel A-Q, PG Apartments).",
                "Transport preferences (Shared Cab, Shared Auto, or Either).",
                "Highlights that students don't need exact same-minute departures."
            ],
            "image": os.path.join(SCREENSHOT_DIR, "4_journey_form.png")
        },
        {
            "slide_num": "05",
            "tag": "SMART ALGORITHM",
            "title": "Compatibility Ranking & Match Scores",
            "subtitle": "Multi-factor algorithm evaluating time, hostel, & destination",
            "bullets": [
                "Calculates weighted Match Score % (e.g. 96% Match).",
                "Hostel Proximity Scoring: Same hostel > Nearby quad > Campus.",
                "Clear reasoning breakdown badges (✓ Same hostel, ✓ Within 5 mins).",
                "Allows joining existing open groups or pairing with individual students."
            ],
            "image": os.path.join(SCREENSHOT_DIR, "5_matching_results.png")
        },
        {
            "slide_num": "06",
            "tag": "GROUP FORMATION",
            "title": "Group Details & Pickup Point Hubs",
            "subtitle": "Group capacity limits (2-4 students) & meeting spot optimization",
            "bullets": [
                "Dynamic group formation enforcing 4-member maximum capacity.",
                "Automatic suggested pickup point calculation (e.g. Main Gate).",
                "Shows student member cards, departure windows, & contact info.",
                "Journey Confirmation button to finalize ride status among members."
            ],
            "image": os.path.join(SCREENSHOT_DIR, "6_group_details_3.png")
        },
        {
            "slide_num": "07",
            "tag": "DYNAMIC SIMULATOR",
            "title": "Real-Time Student Join Simulator",
            "subtitle": "Interactive demo of dynamic matching & group capacity limits",
            "bullets": [
                "Interactive simulator button adds incoming students (Simran Kaur • H4).",
                "Demonstrates live transition from 3/4 members to 4/4 Full capacity.",
                "Capacity alert enforcement automatically routes extra students to Group #105.",
                "Allows one-click reset to test dynamic matching logic multiple times."
            ],
            "image": os.path.join(SCREENSHOT_DIR, "7_group_details_4.png")
        },
        {
            "slide_num": "08",
            "tag": "GROUP CHAT",
            "title": "Student Group Coordination Chat",
            "subtitle": "Interactive frontend messaging for pickup & fare split",
            "bullets": [
                "In-app instant messaging interface for matched group members.",
                "Pre-populated realistic student coordination conversation.",
                "Supports typing and sending live messages with timestamping.",
                "Encourages student alignment on exact pickup time and cab fare split."
            ],
            "image": os.path.join(SCREENSHOT_DIR, "8_group_chat.png")
        },
        {
            "slide_num": "09",
            "tag": "JOURNEY HISTORY",
            "title": "My Journeys & Trip Logs",
            "subtitle": "Tabbed view of active upcoming & completed TIET trips",
            "bullets": [
                "Upcoming Journeys tab with active group status & departure times.",
                "Completed Journeys log detailing date, destination, & travel partners.",
                "Student reliability rating system (4.8/5.0 stars) for trust & safety.",
                "Quick access to group chats and trip confirmation details."
            ],
            "image": os.path.join(SCREENSHOT_DIR, "9_my_journeys.png")
        },
        {
            "slide_num": "10",
            "tag": "STUDENT PROFILE",
            "title": "Student Profile & Platform Impact",
            "subtitle": "Verified student identity & project summary",
            "bullets": [
                "Verified profile for Kartik Arora (Hostel H3 • TIET CoE).",
                "Tracks total trips completed (12) & reliability rating (4.8★).",
                "Impact: Cost reduction, safety enhancement, & community building.",
                "Frontend Prototype built using React.js, Tailwind CSS, & Vite."
            ],
            "image": os.path.join(SCREENSHOT_DIR, "10_profile.png")
        }
    ]

    blank_layout = prs.slide_layouts[6]

    for data in slides_data:
        slide = prs.slides.add_slide(blank_layout)

        # Header Bar Background
        header_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.3))
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = NAVY
        header_shape.line.fill.background()

        # Tag Pill (Enlarged)
        txBox = slide.shapes.add_textbox(Inches(0.6), Inches(0.12), Inches(6), Inches(0.35))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = f"SLIDE {data['slide_num']} • {data['tag']}"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = RGBColor(147, 197, 253)

        # Slide Title (Enlarged)
        txBox_title = slide.shapes.add_textbox(Inches(0.6), Inches(0.45), Inches(12), Inches(0.75))
        tf_title = txBox_title.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = data['title']
        p_title.font.size = Pt(25)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(255, 255, 255)

        # Left Column: Subtitle & Bullet Points (ENLARGED FONT SIZES)
        left_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(5.8), Inches(5.6))
        tf_left = left_box.text_frame
        tf_left.word_wrap = True

        # Subtitle (Enlarged to Pt 18)
        p_sub = tf_left.paragraphs[0]
        p_sub.text = data['subtitle']
        p_sub.font.size = Pt(18)
        p_sub.font.bold = True
        p_sub.font.color.rgb = BLUE
        p_sub.space_after = Pt(16)

        # Bullet Points (Enlarged to Pt 15)
        for bullet in data['bullets']:
            p_b = tf_left.add_paragraph()
            p_b.text = f"•  {bullet}"
            p_b.font.size = Pt(15)
            p_b.font.bold = False
            p_b.font.color.rgb = TEXT_DARK
            p_b.space_after = Pt(12)

        # Right Column: Screenshot Image
        if os.path.exists(data['image']):
            img_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), Inches(1.5), Inches(6.2), Inches(5.4))
            img_box.fill.solid()
            img_box.fill.fore_color.rgb = ACCENT_BG
            img_box.line.color.rgb = RGBColor(203, 213, 225)
            
            slide.shapes.add_picture(data['image'], Inches(6.7), Inches(1.6), width=Inches(6.0))

    ppt_path = r"C:\Users\arora\.gemini\antigravity\scratch\thapar-journey-connect\Thapar_Journey_Connect_Presentation.pptx"
    try:
        prs.save(ppt_path)
        print(f"Successfully saved PowerPoint presentation to {ppt_path}")
    except PermissionError:
        ppt_path_v2 = r"C:\Users\arora\.gemini\antigravity\scratch\thapar-journey-connect\Thapar_Journey_Connect_Presentation_Enlarged.pptx"
        prs.save(ppt_path_v2)
        print(f"Original PPT file was locked; saved updated presentation to {ppt_path_v2}")

if __name__ == "__main__":
    images_exist = all(
        os.path.exists(os.path.join(SCREENSHOT_DIR, f"{i}_{name}.png"))
        for i, name in [
            (1, "auth_page"), (2, "otp_screen"), (3, "dashboard"), (4, "journey_form"),
            (5, "matching_results"), (6, "group_details_3"), (7, "group_details_4"),
            (8, "group_chat"), (9, "my_journeys"), (10, "profile")
        ]
    )
    if not images_exist:
        asyncio.run(capture_screenshots())
    build_powerpoint()
